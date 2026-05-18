"""Update Section 12 in UAP_Architecture_Document_formatted.docx:
  1. Refine 12.1 Onboarding Agent intro to make the three-phase boundary explicit
  2. Update Figure 12.1 caption to match the new sequence diagram
  3. Replace embedded image1.png (rId11 → /word/media/image1.png) with the
     newly rendered docs/architecture/sequence-onboarding-agent.png
"""
import os
import shutil
import zipfile
from docx import Document
from docx.oxml.ns import qn

DOC = "UAP_Architecture_Document_formatted.docx"
NEW_IMG = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "architecture", "sequence-onboarding-agent.png")
)
assert os.path.exists(NEW_IMG), f"missing {NEW_IMG}"

# --- Step 1 + 2: text edits via python-docx -------------------------------
doc = Document(DOC)

# Locate paragraphs in 12.1
p_intro = None       # the long intro after "12.1 Onboarding Agent"
p_caption = None     # "Figure 12.1 ..."
in_121 = False
for p in doc.paragraphs:
    t = p.text
    if t.strip().startswith("12.1 Onboarding Agent"):
        in_121 = True
        continue
    if in_121:
        if t.strip().startswith("12.2"):
            break
        if p_intro is None and t.startswith("The Onboarding Agent automates"):
            p_intro = p
        if t.startswith("Figure 12.1"):
            p_caption = p

assert p_intro is not None, "could not find 12.1 intro paragraph"
assert p_caption is not None, "could not find Figure 12.1 caption"


def replace_paragraph_text(paragraph, new_text):
    """Replace text of a paragraph while preserving its first run's formatting."""
    runs = paragraph.runs
    if not runs:
        run = paragraph.add_run(new_text)
        return
    # Capture formatting from first run
    first = runs[0]
    font_name = first.font.name
    font_size = first.font.size
    bold = first.bold
    italic = first.italic
    color = first.font.color.rgb if first.font.color and first.font.color.type is not None else None
    # Remove all existing runs
    for r in list(runs):
        r._element.getparent().remove(r._element)
    for r in paragraph._element.findall(qn("w:r")):
        paragraph._element.remove(r)
    new_run = paragraph.add_run(new_text)
    if font_name:
        new_run.font.name = font_name
    if font_size:
        new_run.font.size = font_size
    if bold is not None:
        new_run.bold = bold
    if italic is not None:
        new_run.italic = italic
    if color is not None:
        new_run.font.color.rgb = color


new_intro = (
    "The Onboarding Agent automates the seven-stage publish-to-production pipeline that takes "
    "an Agent, MCP server or Tool from a developer submission to a live, governed endpoint in "
    "hours rather than weeks. It operates around the developer's source-control and CI/CD "
    "boundary in three explicit phases. "
    "Phase 1 — Pre-push (co-pilot): in the IDE / Publisher Portal the agent scaffolds the "
    "AgentCard / MCP manifest, infers capabilities and data sensitivity from source, lints "
    "against tenant policy, generates eval stubs and opens the initial PR with the right "
    "labels and CODEOWNERS. "
    "Phase 2 — Push and CI (developer-driven): the developer's push (or the resulting PR) is "
    "what triggers GitHub Actions; the agent does not invoke CI directly. The hardened builder "
    "performs the build, dependency audit and security / supply-chain scan, emits the SBOM "
    "(CycloneDX) and the SLSA-style provenance attestation (Sigstore-signed via OIDC), and "
    "runs the eval harness — publishing all outputs as workflow artifacts. "
    "Phase 3 — Post-push (verify, govern, deploy): the agent ingests those Actions outputs "
    "via webhook / Event Grid, cryptographically verifies the SLSA attestation against "
    "marketplace policy, registers the AgentCard in MongoDB with status=staging, assembles a "
    "human-actionable governance ticket (provenance summary, SBOM / CVE highlights, eval delta, "
    "suggested approver), drives Responsible AI / Security approval gates for sensitive or "
    "production-tier assets, coordinates infrastructure provisioning on Azure Container Apps "
    "with Managed Identity and APIM routing, and oversees the controlled canary rollout "
    "(5% → 25% → 50% → 100%)."
)
replace_paragraph_text(p_intro, new_intro)

new_caption = (
    "Figure 12.1 — Onboarding Agent sequence: three explicit phases — "
    "Phase 1 pre-push co-pilot (scaffold / lint / open PR), "
    "Phase 2 developer-driven push that triggers GitHub Actions (build, scan, SBOM, "
    "SLSA attestation, eval — agent does not invoke CI), "
    "Phase 3 post-push verify-govern-deploy (SLSA verification, registry, governance ticket, "
    "approval, provisioning, canary)."
)
replace_paragraph_text(p_caption, new_caption)

doc.save(DOC)
print("docx text updated")

# --- Step 3: swap /word/media/image1.png -----------------------------------
# Re-open the docx (a zip) and replace the embedded image binary.
tmp = DOC + ".tmp"
with zipfile.ZipFile(DOC, "r") as zin:
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/media/image1.png":
                with open(NEW_IMG, "rb") as f:
                    data = f.read()
                print(f"replaced {item.filename} ({len(data)} bytes)")
            zout.writestr(item, data)
shutil.move(tmp, DOC)
print("docx image swapped, saved", DOC)
