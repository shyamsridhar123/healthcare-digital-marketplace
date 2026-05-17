"""Generate a PPTX slide deck from sequence-diagrams.html.

Parses the HTML to extract titles, subtitles, tags, and Mermaid source,
then renders each diagram to PNG via mermaid-cli (mmdc / npx) and
assembles them into a dark-themed PowerPoint presentation.

Usage:
    python generate_pptx.py
"""

import html
import json
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# ── Paths ────────────────────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).resolve().parent
HTML_FILE = SCRIPT_DIR / "sequence-diagrams.html"
OUTPUT_PPTX = SCRIPT_DIR / "sequence-diagrams.pptx"

# ── Theme colours ────────────────────────────────────────────────────
BG = RGBColor(0x0F, 0x17, 0x2A)
CARD = RGBColor(0x1A, 0x23, 0x3B)
ACCENT = RGBColor(0x00, 0x78, 0xD4)
TEAL = RGBColor(0x00, 0xE5, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xA0, 0xAE, 0xC0)
TAG_CLR = RGBColor(0x8B, 0x5C, 0xF6)

# ── Mermaid theme for PNG render ─────────────────────────────────────
MERMAID_CONFIG = {
    "theme": "dark",
    "themeVariables": {
        "primaryColor": "#1a233b",
        "primaryBorderColor": "#0078d4",
        "primaryTextColor": "#e0e8f0",
        "secondaryColor": "#303d55",
        "tertiaryColor": "#0f172a",
        "lineColor": "#00e5ff",
        "textColor": "#e0e8f0",
        "actorBkg": "#0078d4",
        "actorBorder": "#0078d4",
        "actorTextColor": "#ffffff",
        "activationBorderColor": "#00e5ff",
        "activationBkgColor": "#1a233b",
        "sequenceNumberColor": "#000000",
        "signalColor": "#00e5ff",
        "signalTextColor": "#ffffff",
        "noteBkgColor": "#303d55",
        "noteTextColor": "#f5e6c2",
        "noteBorderColor": "#f59e0b",
        "labelBoxBkgColor": "#1e293b",
        "labelBoxBorderColor": "#0078d4",
        "labelTextColor": "#e0e8f0",
    },
    "sequence": {
        "diagramMarginX": 30,
        "diagramMarginY": 15,
        "actorMargin": 60,
        "width": 180,
        "height": 50,
        "boxMargin": 6,
        "boxTextMargin": 8,
        "noteMargin": 12,
        "messageMargin": 40,
        "mirrorActors": True,
        "useMaxWidth": False,
        "wrap": True,
        "wrapPadding": 15,
    },
    "fontFamily": "Segoe UI, system-ui, sans-serif",
    "fontSize": 13,
}


# ── 1. Parse HTML ────────────────────────────────────────────────────

def parse_slides(html_text: str):
    """Return list of dicts: {title, subtitle, tag, mermaid}."""
    pattern = re.compile(
        r'<div\s+class="slide"'
        r'\s+data-title="([^"]*)"'
        r'\s+data-subtitle="([^"]*)"'
        r'\s+data-tag="([^"]*)"'
        r'[^>]*>'
        r'.*?<pre\s+class="mermaid">(.*?)</pre>',
        re.DOTALL,
    )
    slides = []
    for m in pattern.finditer(html_text):
        slides.append({
            "title": html.unescape(m.group(1)),
            "subtitle": html.unescape(m.group(2)),
            "tag": html.unescape(m.group(3)),
            "mermaid": html.unescape(m.group(4)).strip(),
        })
    return slides


# ── 2. Render Mermaid → PNG ──────────────────────────────────────────

def render_mermaid_to_png(mermaid_code: str, out_png: Path, config_path: Path):
    """Render mermaid code to a PNG file using npx mmdc."""
    mmd_file = out_png.with_suffix(".mmd")
    mmd_file.write_text(mermaid_code, encoding="utf-8")
    cmd = [
        "npx", "--yes", "@mermaid-js/mermaid-cli",
        "-i", str(mmd_file),
        "-o", str(out_png),
        "-c", str(config_path),
        "-b", "#0f172a",
        "-w", "1600",
        "--scale", "2",
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=120, shell=True)
    if result.returncode != 0:
        print(f"  WARN mmdc failed: {result.stderr[:300]}")
        return False
    return out_png.exists()


# ── 3. Build PPTX ───────────────────────────────────────────────────

def set_slide_bg(slide, color: RGBColor):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Segoe UI"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def build_title_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, BG)

    # Accent line
    from pptx.shapes.autoshape import Shape
    shape = slide.shapes.add_shape(
        1, Inches(0), Inches(3.3), Inches(13.33), Inches(0.06)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()

    add_text_box(slide, Inches(1), Inches(1.5), Inches(11), Inches(1.2),
                 "AI Platform — Sequence Diagrams", font_size=36, color=WHITE, bold=True,
                 alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(2.5), Inches(11), Inches(0.6),
                 "Developer Experience & RCM Domain Orchestration Flows",
                 font_size=18, color=MUTED, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(4.0), Inches(11), Inches(0.5),
                 "12 Sequence Diagrams  •  6 Dev-Experience  •  6 RCM Domain Workflows",
                 font_size=14, color=TEAL, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(6.5), Inches(11), Inches(0.4),
                 "Microsoft Agent Framework  •  Azure Container Apps  •  Cosmos DB  •  Azure OpenAI",
                 font_size=11, color=MUTED, alignment=PP_ALIGN.CENTER)


def build_diagram_slide(prs: Presentation, idx: int, slide_data: dict, png_path: Path | None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, BG)

    # ── Badge + Title row
    badge_box = add_text_box(slide, Inches(0.4), Inches(0.25), Inches(0.5), Inches(0.4),
                             str(idx), font_size=14, color=WHITE, bold=True,
                             alignment=PP_ALIGN.CENTER)
    # Colour the badge background
    badge_box.fill.solid()
    badge_box.fill.fore_color.rgb = ACCENT

    add_text_box(slide, Inches(1.0), Inches(0.25), Inches(9), Inches(0.4),
                 slide_data["title"], font_size=22, color=WHITE, bold=True)

    # ── Subtitle
    add_text_box(slide, Inches(1.0), Inches(0.65), Inches(8), Inches(0.35),
                 slide_data["subtitle"], font_size=11, color=MUTED)

    # ── Tag pill
    tag_box = add_text_box(slide, Inches(10.0), Inches(0.25), Inches(3), Inches(0.35),
                           slide_data["tag"], font_size=9, color=TAG_CLR, bold=True,
                           alignment=PP_ALIGN.RIGHT)

    # ── Accent line below header
    shape = slide.shapes.add_shape(
        1, Inches(0.4), Inches(1.05), Inches(12.5), Inches(0.03)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()

    # ── Diagram image (or text fallback)
    if png_path and png_path.exists():
        # Calculate placement to fit within the card area
        from PIL import Image
        img = Image.open(png_path)
        img_w, img_h = img.size
        img.close()

        max_w = Inches(12.2)
        max_h = Inches(5.8)
        card_left = Inches(0.55)
        card_top = Inches(1.25)

        # Scale to fit
        ratio = min(max_w / Emu(int(img_w * 914400 / 96)),
                    max_h / Emu(int(img_h * 914400 / 96)))
        if ratio < 1:
            pic_w = int(img_w * 914400 / 96 * ratio)
            pic_h = int(img_h * 914400 / 96 * ratio)
        else:
            pic_w = int(img_w * 914400 / 96)
            pic_h = int(img_h * 914400 / 96)

        # Centre horizontally
        pic_left = card_left + (max_w - pic_w) // 2
        pic_top = card_top + (max_h - pic_h) // 2

        slide.shapes.add_picture(str(png_path), pic_left, pic_top, pic_w, pic_h)
    else:
        # Fallback: show mermaid source as text
        code = slide_data["mermaid"]
        # Truncate if very long
        if len(code) > 2000:
            code = code[:2000] + "\n... (truncated)"
        card = slide.shapes.add_shape(
            1, Inches(0.5), Inches(1.2), Inches(12.3), Inches(5.9)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.color.rgb = RGBColor(0x30, 0x3D, 0x55)

        add_text_box(slide, Inches(0.7), Inches(1.4), Inches(11.9), Inches(5.5),
                     code, font_size=8, color=RGBColor(0xE0, 0xE8, 0xF0),
                     font_name="Consolas")


def build_pptx(slides_data: list, png_dir: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    build_title_slide(prs)

    for i, sd in enumerate(slides_data, 1):
        png_path = png_dir / f"slide_{i}.png"
        build_diagram_slide(prs, i, sd, png_path)
        print(f"  Built slide {i}: {sd['title']}")

    prs.save(str(OUTPUT_PPTX))
    print(f"\n✅ Saved: {OUTPUT_PPTX}")


# ── Main ─────────────────────────────────────────────────────────────

def main():
    print("Parsing HTML...")
    html_text = HTML_FILE.read_text(encoding="utf-8")
    slides_data = parse_slides(html_text)
    print(f"  Found {len(slides_data)} slides")

    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir = Path(tmpdir)

        # Write mermaid config
        config_path = tmpdir / "mermaid-config.json"
        config_path.write_text(json.dumps(MERMAID_CONFIG), encoding="utf-8")

        # Render each diagram
        print("\nRendering diagrams to PNG (this may take a minute)...")
        for i, sd in enumerate(slides_data, 1):
            png_path = tmpdir / f"slide_{i}.png"
            print(f"  [{i}/{len(slides_data)}] {sd['title']}...", end=" ", flush=True)
            ok = render_mermaid_to_png(sd["mermaid"], png_path, config_path)
            print("✓" if ok else "✗ (will use text fallback)")

        # Build PPTX
        print("\nBuilding PPTX...")
        build_pptx(slides_data, tmpdir)


if __name__ == "__main__":
    main()
