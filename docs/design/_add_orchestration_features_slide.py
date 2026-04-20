"""Append an 'Orchestration Design Features' slide to the deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
import os

SRC = r'c:\gitrepos\ai-marketplace\docs\design\Global_Agentic_Orchestrator_Design_v2_with_decisions.pptx'
OUT = r'c:\gitrepos\ai-marketplace\docs\design\Global_Agentic_Orchestrator_Design_v2_with_decisions.pptx'

# If locked, write to a new file
try:
    with open(OUT, 'ab'):
        pass
except PermissionError:
    OUT = SRC.replace('.pptx', '_v3.pptx')

prs = Presentation(SRC)
SW, SH = prs.slide_width, prs.slide_height

NAVY = RGBColor(0x0A, 0x2A, 0x5E)
BLUE = RGBColor(0x1F, 0x5F, 0xB8)
TEAL = RGBColor(0x0F, 0x8A, 0x8A)
ORANGE = RGBColor(0xE3, 0x6A, 0x1F)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
PURPLE = RGBColor(0x6A, 0x3F, 0xA0)
RED = RGBColor(0xB5, 0x3A, 0x3A)
LIGHT = RGBColor(0xF4, 0xF7, 0xFB)
LIGHTER = RGBColor(0xFA, 0xFB, 0xFD)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x5E, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
slide = prs.slides.add_slide(blank_layout)


def add_text(slide, x, y, w, h, text, *, size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT, fill=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    if fill is not None:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fill
        tb.line.fill.background()
    tf = tb.text_frame
    tf.margin_left = Emu(60000)
    tf.margin_right = Emu(60000)
    tf.margin_top = Emu(40000)
    tf.margin_bottom = Emu(40000)
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = 'Segoe UI'
    return tb


def add_rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


# Background
add_rect(slide, 0, 0, SW, SH, WHITE)

# Header
add_rect(slide, 0, 0, SW, Inches(0.95), NAVY)
add_text(slide, Inches(0.4), Inches(0.12), SW - Inches(0.8), Inches(0.5),
         'Orchestration Design Features', size=24, bold=True, color=WHITE)
add_text(slide, Inches(0.4), Inches(0.55), SW - Inches(0.8), Inches(0.35),
         'Core capabilities the Global Agentic Orchestrator delivers out of the box',
         size=12, color=RGBColor(0xCF, 0xDC, 0xF0))

# Feature cards: 4 cols x 3 rows
rows, cols = 3, 4
margin_x = Inches(0.3)
margin_top = Inches(1.15)
margin_bottom = Inches(0.55)
gap = Inches(0.15)

card_w = (SW - margin_x * 2 - gap * (cols - 1)) / cols
card_h = (SH - margin_top - margin_bottom - gap * (rows - 1)) / rows

features = [
    ('Graph-Based Workflows', BLUE,
     'LangGraph nodes + edges with conditional branching, parallel fan-out, loops, and sub-graphs for complex RCM processes.'),
    ('Durable State & Checkpoints', TEAL,
     'Every step persisted to MongoDB. Resume after hours or days for human approvals — zero thread blocking.'),
    ('Dynamic Skill Resolution', GREEN,
     'Capability-tag matching against registry with confidence scoring. Orchestrator picks best agent at runtime.'),
    ('Declarative Plan Templates', ORANGE,
     'YAML/JSON blueprints for recurring workflows (prior-auth, denials, eligibility). Version-controlled and shareable.'),
    ('Human-in-the-Loop Gates', PURPLE,
     'Native interrupts for review, edit, or approve. State mutations allowed mid-flow without losing context.'),
    ('Parallel Fan-Out / Fan-In', BLUE,
     'Invoke N agents concurrently; aggregate with merge strategies (consensus, first-success, weighted vote).'),
    ('Policy & Guardrail Hooks', RED,
     'Pre/post-execution policy evaluation: PHI/PII redaction, content filters, cost caps, RBAC checks.'),
    ('Retry, Timeout & Compensation', TEAL,
     'Per-node retry policies with backoff. Saga-style compensation for rollback of partially completed workflows.'),
    ('Streaming & Partial Results', GREEN,
     'Token-level streaming to UI. Intermediate node outputs surfaced live via SSE/WebSocket.'),
    ('Multi-Modal Context Window', ORANGE,
     'Unified memory: Redis session + Mongo long-term + vector RAG. Context assembled per agent invocation.'),
    ('End-to-End OTel Tracing', PURPLE,
     'Trace IDs propagated across agents, tools, and models. Unified timeline in App Insights + Arize Phoenix.'),
    ('Cost & Token Governance', RED,
     'Per-tenant budgets, per-workflow token caps, model routing (GPT-5 → GPT-4o → 4o-mini) by task complexity.'),
]

for idx, (title, accent, body) in enumerate(features):
    r = idx // cols
    c = idx % cols
    x = margin_x + c * (card_w + gap)
    y = margin_top + r * (card_h + gap)

    add_rect(slide, x, y, card_w, card_h, LIGHTER)
    add_rect(slide, x, y, Inches(0.12), card_h, accent)
    add_text(slide, x + Inches(0.25), y + Inches(0.15), card_w - Inches(0.35), Inches(0.45),
             title, size=12, bold=True, color=NAVY)
    add_text(slide, x + Inches(0.25), y + Inches(0.65), card_w - Inches(0.35), card_h - Inches(0.75),
             body, size=9.5, color=GREY)

# Footer
footer_y = SH - Inches(0.48)
add_rect(slide, 0, footer_y, SW, Inches(0.48), NAVY)
add_text(slide, Inches(0.4), footer_y + Inches(0.09), SW - Inches(0.8), Inches(0.35),
         'All features available via a single declarative workflow spec  •  Framework-agnostic A2A + MCP  •  HIPAA-compliant by default',
         size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(f'Saved to {OUT}. Slide count: {len(prs.slides)}')
