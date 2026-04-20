from pptx import Presentation
import sys

p = Presentation(r'c:\gitrepos\ai-marketplace\docs\design\Global_Agentic_Orchestrator_Design_v2.pptx')
for i, s in enumerate(p.slides):
    title = ''
    try:
        if s.shapes.title:
            title = s.shapes.title.text
    except Exception:
        pass
    print(f'--- Slide {i+1}: {title} ---')
    for sh in s.shapes:
        if sh.has_text_frame:
            txt = sh.text_frame.text
            if txt.strip():
                print(txt)
    print()
