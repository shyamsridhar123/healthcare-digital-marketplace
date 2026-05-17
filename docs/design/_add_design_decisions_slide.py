"""Append a 'Key Design Decisions' summary slide to Global_Agentic_Orchestrator_Design_v2.pptx."""
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

SRC = r'c:\gitrepos\ai-marketplace\docs\design\Global_Agentic_Orchestrator_Design_v2.pptx'
OUT = r'c:\gitrepos\ai-marketplace\docs\design\Global_Agentic_Orchestrator_Design_v2_with_decisions.pptx'

prs = Presentation(SRC)
SW, SH = prs.slide_width, prs.slide_height

# Palette (matches Optum / Microsoft blue accents used elsewhere in the deck)
NAVY = RGBColor(0x0A, 0x2A, 0x5E)
BLUE = RGBColor(0x1F, 0x5F, 0xB8)
TEAL = RGBColor(0x0F, 0x8A, 0x8A)
ORANGE = RGBColor(0xE3, 0x6A, 0x1F)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
LIGHT = RGBColor(0xF4, 0xF7, 0xFB)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x5E, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Use a blank layout
blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
slide = prs.slides.add_slide(blank_layout)


def add_text(slide, x, y, w, h, text, *, size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT, fill=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    if fill is not None:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fill
        tb.line.fill.background()
    tf = tb.text_frame
    tf.margin_left = Emu(60000)
    tf.margin_right = Emu(60000)
    tf.margin_top = Emu(40000)
    tf.margin_bottom = Emu(40000)
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = 'Segoe UI'
    return tb


def add_rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp


# Background
add_rect(slide, 0, 0, SW, SH, WHITE)

# Header bar
add_rect(slide, 0, 0, SW, Inches(0.95), NAVY)
add_text(slide, Inches(0.4), Inches(0.12), SW - Inches(0.8), Inches(0.5),
         'Key Design Decisions — Summary', size=24, bold=True, color=WHITE)
add_text(slide, Inches(0.4), Inches(0.55), SW - Inches(0.8), Inches(0.35),
         'Ten choices that define the Global Agentic Orchestrator for RCM',
         size=12, color=RGBColor(0xCF, 0xDC, 0xF0))

# Layout: 5 columns x 2 rows of decision cards
rows, cols = 2, 5
margin_x = Inches(0.3)
margin_top = Inches(1.15)
margin_bottom = Inches(0.55)
gap = Inches(0.15)

card_w = (SW - margin_x * 2 - gap * (cols - 1)) / cols
card_h = (SH - margin_top - margin_bottom - gap * (rows - 1)) / rows

decisions = [
    ('01', 'Two-Tier Orchestration', BLUE,
     'Central Global Orchestrator + specialized Domain Orchestrators. Separates cross-domain policy from in-domain expertise.'),
    ('02', 'LangGraph over MAF', TEAL,
     'Mature graph engine with durable checkpoints and framework-agnostic A2A. MAF strengths rebuilt as platform capabilities.'),
    ('03', 'Direct A2A Calls (Default)', GREEN,
     'Sync API, async behavior via LangGraph durable pauses. Lowest latency; Event Hub added only for bursty fan-out.'),
    ('04', 'Capability-Based Discovery', ORANGE,
     'AgentCards in Mongo registry + confidence-scored dynamic resolver. Cross-domain reuse, no hardcoded endpoints.'),
    ('05', 'MCP for Enterprise Tools', BLUE,
     'All enterprise APIs (claims, EHR, billing) wrapped as MCP servers behind APIM. Uniform tool surface; zero-trust via Entra ID.'),
    ('06', 'Three-Plane Architecture', TEAL,
     'Model Plane (UAIS/Azure OpenAI), Tool Plane (MCP), Agent Plane (containers). Unified memory: Redis + Mongo + vector index.'),
    ('07', 'Automated 7-Stage Onboarding', GREEN,
     'Submission → Validation → Registration → Approval → Provisioning → Test → Canary. Weeks to hours with human gates where needed.'),
    ('08', 'Canary + Instant Rollback', ORANGE,
     'Semantic versioning with 5 → 25 → 50 → 100% traffic shifts. ACA revisions enable sub-minute revert on regression.'),
    ('09', 'Arize Phoenix LLM Ops', BLUE,
     'Self-hosted on Azure for HIPAA. OTel traces + LLM-as-a-judge + drift detection correlated with Azure Monitor infra metrics.'),
    ('10', 'Compliance by Declaration', TEAL,
     'AgentCards carry Responsible AI + data-category tags. Security Agent enforces; PHI/PII redaction and audit trail by design.'),
]

for idx, (num, title, accent, body) in enumerate(decisions):
    r = idx // cols
    c = idx % cols
    x = margin_x + c * (card_w + gap)
    y = margin_top + r * (card_h + gap)

    # Card background
    add_rect(slide, x, y, card_w, card_h, LIGHT)
    # Accent stripe
    add_rect(slide, x, y, card_w, Inches(0.22), accent)
    # Number badge
    add_text(slide, x + Inches(0.1), y + Inches(0.02), Inches(0.5), Inches(0.2),
             num, size=11, bold=True, color=WHITE)
    # Title
    add_text(slide, x + Inches(0.1), y + Inches(0.3), card_w - Inches(0.2), Inches(0.5),
             title, size=12, bold=True, color=NAVY)
    # Body
    add_text(slide, x + Inches(0.1), y + Inches(0.75), card_w - Inches(0.2), card_h - Inches(0.85),
             body, size=9, color=GREY)

# Footer takeaway
footer_y = SH - Inches(0.48)
add_rect(slide, 0, footer_y, SW, Inches(0.48), NAVY)
add_text(slide, Inches(0.4), footer_y + Inches(0.09), SW - Inches(0.8), Inches(0.35),
         'Net effect:  faster onboarding (weeks → hours)  •  40–50% agent reuse  •  HIPAA-ready by default  •  scale to 100K+ encounters/day',
         size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(f'Saved to {OUT}. Slide count: {len(prs.slides)}')
